"""
Turn the Phase-2 slide PDF into a PowerPoint file (one full-slide picture per page), for presenting from
PowerPoint / Google Slides. Needs python-pptx and poppler's pdftoppm.
Usage: python scripts/slides_to_pptx.py
"""
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
PDF = ROOT / 'docs' / 'presentation' / 'IndicDoc_AI_Phase2_Implementation.pdf'


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(['pdftoppm', '-r', '150', '-png', str(PDF), f'{tmp}/s'], check=True)
        for png in sorted(Path(tmp).glob('s-*.png')):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
    out = PDF.with_suffix('.pptx')
    prs.save(out)
    print('Wrote', out)


if __name__ == '__main__':
    main()
