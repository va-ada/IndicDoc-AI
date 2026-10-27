"""
IndicDoc AI - Phase-2 deck in the department's template format: black and white, portrait US-Letter pages,
centred serif titles, bullet lists and plain bordered tables, the same 18 sections in the same order.

Content and numbers are the same as the colour deck (imported from build_slides.py, which reads outputs/metrics);
figures are converted to greyscale. Writes an editable PPTX and a PDF (LibreOffice).

Usage: python scripts/build_slides_bw.py
"""
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_slides as C  # noqa: E402  (content and measured numbers of the Phase-2 deck)

OUT = C.OUT / 'IndicDoc_AI_Phase2_Implementation_BW.pptx'
FONT, MONO = 'Times New Roman', 'Courier New'
BLACK, WHITE = RGBColor(0, 0, 0), RGBColor(255, 255, 255)
W, H = 8.5, 11.0
L, R = 0.6, 0.6
CW = W - L - R
TMP = Path(tempfile.mkdtemp())


def f3(v):
    return f'{v:.3f}'


def pct(v):
    return f'{100 * v:.0f}%'


def strip(s: str) -> str:
    return re.sub(r'<[^>]+>', '', s).replace('&amp;', '&')


def runs(par, text, size, bold=False, font=FONT):
    """Text with **bold** spans."""
    for k, part in enumerate(re.split(r'\*\*', text)):
        if not part:
            continue
        r = par.add_run()
        r.text = part
        r.font.size, r.font.name, r.font.color.rgb = Pt(size), font, BLACK
        r.font.bold = bold or (k % 2 == 1)


def title(s, text):
    tb = s.shapes.add_textbox(Inches(L), Inches(0.35), Inches(CW), Inches(0.6))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    runs(p, text, 24)
    return tb


def bullets(s, items, top, height, size=13, left=L, width=CW, gap=3):
    """items: str, (str, level), '' for an empty line, or ('#', str) for a numbered list item."""
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        level, text, numbered = 0, it, None
        if isinstance(it, tuple):
            if it[0] == '#':
                numbered, text = it[1], it[2]
            else:
                text, level = it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        if text == '':
            runs(p, ' ', size)
            continue
        sz = size - 1.5 * level
        indent = 0.25 + 0.3 * level
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(int(Inches(indent))))
        pPr.set('indent', str(int(-Inches(0.22 if numbered is None else 0.32))))
        if numbered is None:
            bu = pPr.makeelement(qn('a:buChar'), {'char': '•'})
        else:
            bu = pPr.makeelement(qn('a:buAutoNum'), {'type': 'arabicPeriod', 'startAt': str(numbered)})
        pPr.append(bu)
        runs(p, text, sz)
    return tb


def _border(cell, w=Pt(0.75)):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        ln = tcPr.makeelement(qn(tag), {'w': str(int(w)), 'cmpd': 'sng'})
        fill = ln.makeelement(qn('a:solidFill'), {})
        clr = fill.makeelement(qn('a:srgbClr'), {'val': '000000'})
        fill.append(clr)
        ln.append(fill)
        tcPr.append(ln)


def table(s, head, rows, top, widths, size=10, left=L, row_h=0.28, bold_first_col=False):
    n_r, n_c = len(rows) + 1, len(head)
    shape = s.shapes.add_table(n_r, n_c, Inches(left), Inches(top), Inches(sum(widths)), Inches(row_h * n_r))
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(el)
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i in range(n_r):
        tbl.rows[i].height = Inches(row_h)
        for j in range(n_c):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            _border(cell)
            cell.margin_left = cell.margin_right = Inches(0.04)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            text = head[j] if i == 0 else str(rows[i - 1][j])
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
            runs(p, text, size, bold=(i == 0) or (bold_first_col and j == 0))
    return shape


def grey(path: Path) -> Path:
    out = TMP / (path.stem + '_bw.png')
    if not out.exists():
        im = Image.open(path).convert('L')
        im.save(out)
    return out


def image(s, path, left, top, width=None, height=None, frame=True):
    p = s.shapes.add_picture(str(grey(Path(path))), Inches(left), Inches(top),
                             Inches(width) if width else None, Inches(height) if height else None)
    if frame:
        p.line.color.rgb = BLACK
        p.line.width = Pt(0.5)
    return p


def figure(s, path, left, top, width=None, height=None, cap='', cap_size=9.5):
    """Picture with its caption directly below; returns the bottom edge in inches."""
    pic = image(s, path, left, top, width, height)
    bottom = (pic.top + pic.height) / 914400
    if cap:
        caption(s, cap, pic.left / 914400 - 0.1, bottom + 0.03, pic.width / 914400 + 0.2, size=cap_size)
        bottom += 0.35
    return bottom


def caption(s, text, left, top, width, size=9.5, italic=True):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.3))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    runs(p, text, size)
    for r in p.runs:
        r.font.italic = italic
    return tb


def box(s, x, y, w, h, head, sub='', size=10.5, fill=WHITE):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = BLACK
    sh.line.width = Pt(1)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    runs(p, head, size, bold=True)
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        runs(p2, sub, size - 2)
    return sh


def arrow(s, x1, y1, x2, y2):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = BLACK
    c.line.width = Pt(1.25)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    return c


def code_box(s, text, x, y, w, h, size=9):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = WHITE
    sh.line.color.rgb = BLACK
    sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for k, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        runs(p, line or ' ', size, font=MONO)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    new = lambda: prs.slides.add_slide(blank)  # noqa: E731
    A, P, B, V, VN = C.A, C.P, C.B, C.V, C.VN
    QAP, cq, dq = C.QAP, C.clean_q, C.deg_q
    q_ap = C.q_ap

    # 1. Title slide
    s = new()
    title(s, 'IndicDoc Verify')
    bullets(s, [
        '**Project title:** IndicDoc Verify: Deep-Learning Layout Analysis for Pre-Print Checking of Indian Question '
        'Papers and Their Translations',
        '**Team members:**',
        *[(f'{n} (Roll No. {r})', 1) for _, n, r in sorted(C.TEAM, key=lambda t: t[2])],
        '**Course:** DL Mini Project (Expt. 11, SDG-based)', '',
        'Department of Artificial Intelligence and Machine Learning',
        'St. Francis Institute of Technology (SFIT)', '',
        f'**Project guide:** {C.GUIDE}',
        '**Academic year / semester:** 2026–27, Semester V (TE AIML)',
    ], 1.15, 6, size=14, gap=5)

    # 2. Content
    s = new()
    title(s, 'Content')
    bullets(s, [('#', k, t) for k, t in enumerate(C.items, 1)], 1.15, 7, size=14, gap=2)

    # 3. Literature survey summary table
    s = new()
    title(s, 'Literature Survey Summary Table')
    table(s, ['Ref', 'Author', 'Problem', 'Dataset', 'Method', 'Limitation'],
          [(r, a, p, d, m, l) for r, a, p, d, m, l in C.lit], 1.1, [0.42, 0.9, 1.35, 1.45, 1.5, 1.68], size=8.5,
          row_h=0.5)

    # 4. Research gaps
    s = new()
    title(s, 'Research Gaps')
    table(s, ['#', 'Gap in existing work', 'How this project addresses it'],
          [(c, strip(a), strip(b)) for a, b, c in C.gaps], 1.15, [0.45, 3.4, 3.45], size=11, row_h=0.75)
    bullets(s, ['Court e-filing defect detection already exists (Supreme Court with IIT Madras; a commercial tool '
                'reported in 16 High Courts), so we chose the problem **nobody automates yet**: checking a question '
                'paper and its translations before printing.'], 5.25, 1.5, size=12)

    # 5. Problem statement and objectives
    s = new()
    title(s, 'Problem Statement and Objectives')
    bullets(s, [
        '**Problem:** a question paper is set in one language and translated into Hindi and regional languages; '
        'every version is checked by hand. When a number, option, mark or question differs, lakhs of candidates are '
        'affected (grace marks, re-exams, court cases).',
        '**Main objective:** an offline deep-learning system that reads a printed paper in any of 12 Indian languages '
        'and flags every structural or numerical error, and every difference from its source version, before printing.',
        '**Specific objectives:**',
        ('Detect the paper\'s structure (questions, options, answers) with a layout detector; compare 3 detectors.', 1),
        ('Read text with reliable numbers in 12 languages (number pass).', 1),
        ('Checks: numbering, duplicates, options, marks, translation parity, repeats, format profiles.', 1),
        ('Measure errors found and false alarms on a benchmark; ship a web app and API in Docker.', 1),
        '',
        '**Expected output:** a verdict (ready / check / do not print) and every finding with its question and page, '
        'outlined on the page, as HTML and JSON.',
    ], 1.1, 5.4, size=12.5)
    caption(s, 'It has already happened (news reports)', L, 6.55, CW, size=11, italic=False)
    table(s, ['Exam', 'What went wrong'], C.cases, 6.9, [1.8, 5.5], size=10, row_h=0.36, bold_first_col=True)

    # 6. Proposed system, methodology and novelty
    s = new()
    title(s, 'Proposed System, Methodology and Novelty')
    bullets(s, [
        '**Proposed solution:** IndicDoc Verify reads the paper as it will be printed, finds its structure with a '
        'deep-learning layout detector, reads the text, and checks it alone and against its source version.',
        '**Workflow (input to output):**',
        *[(('#', k, f'**{b}:** {c}')) for k, (_, b, c) in enumerate(C.steps, 1)],
        ('**How deep learning is used:** a convolutional object detector (DocLayout-YOLO) finds every question, '
         f'sub-question, option and answer on the page (AP@50 {min(q_ap):.2f}–{max(q_ap):.2f} on these labels); its '
         'regions turn a page image into structured questions in any script.', 0),
        '',
        '**What is different (novelty):**',
        ('**Translation parity without translation:** numbers, option counts, marks, sub-questions and figures must '
         'match in every language, so versions are compared question by question with no MT model.', 1),
        ('**Number pass:** two OCR models must agree on every digit.', 1),
        ('**USP:** no tool we found checks an Indian question paper against its translations question by question; '
         'this runs offline, in 12 languages, before printing.', 1),
    ], 1.1, 9, size=12.5)

    # 7. System block diagram
    s = new()
    title(s, 'System Block Diagram')
    bullets(s, ['Major components and the flow of data through the system (top to bottom).'], 1.05, 0.5, size=12.5)
    lanes = [('User', [('Streamlit app', 'Verify Paper tab'), ('REST API (FastAPI)', 'POST /verify'),
                       ('Docker, offline', 'papers never leave the exam body')]),
             ('Input', [('Render pages', 'PDF at 200 dpi / images'), ('Source version', 'optional, 12 languages'),
                        ('Question bank, profile', 'earlier papers; JSON rules')]),
             ('Deep learning', [('Layout detector', 'DocLayout-YOLO / YOLOv8s / Faster R-CNN'),
                                ('Post-process', 'high-recall threshold, duplicates removed'),
                                ('Question regions', 'questions, options, answers')]),
             ('Text', [('Indic OCR', 'Tesseract, page\'s own script'), ('Number pass', 'English model, digits agree'),
                       ('Build questions', 'number, stem, options, marks')]),
             ('Checks and output', [('Structure, repeats', 'numbering, duplicates, marks, bank'),
                                    ('Translation parity', 'align versions; compare'),
                                    ('Report', 'verdict, findings on pages; HTML / JSON')])]
    y, bw, bh, x0 = 1.75, 2.0, 0.95, 1.55
    for k, (lane, blocks) in enumerate(lanes):
        lb = s.shapes.add_textbox(Inches(L), Inches(y), Inches(0.95), Inches(bh))
        lb.text_frame.word_wrap = True
        lb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        runs(lb.text_frame.paragraphs[0], lane.upper(), 9.5, bold=True)
        for i, (hd, sub) in enumerate(blocks):
            x = x0 + i * (bw + 0.18)
            box(s, x, y, bw, bh, hd, sub, size=10.5, fill=RGBColor(0xE6, 0xE6, 0xE6) if lane == 'Deep learning' else WHITE)
            if i:
                arrow(s, x - 0.18, y + bh / 2, x, y + bh / 2)
        if k < len(lanes) - 1:
            arrow(s, x0 + 1 * (bw + 0.18) + bw / 2, y + bh, x0 + 1 * (bw + 0.18) + bw / 2, y + bh + 0.45)
        y += bh + 0.45
    caption(s, 'Grey: deep-learning stage. The source version takes the same path as the paper.', L, y + 0.05, CW)

    # 8. Dataset description and pre-processing
    s = new()
    title(s, 'Dataset Description and Pre-processing')
    tot = C.tot
    bullets(s, [
        '**Dataset name and source:** IndicDLP (Nath et al., ICDAR 2025), Hugging Face ai4bharat/indicdlp, AIKosh; '
        '119,806 pages, 12 languages, 12 domains, 42 layout labels.',
        f'**Samples and classes (our subset):** {tot["train"] + tot["val"] + tot["test"]:,} pages, {C.total_boxes:,} '
        f'boxes, 42 classes; {C.qp_pages} question-paper pages, {C.cnt.get("first-level-question", 0):,} question and '
        f'{C.cnt.get("options", 0):,} option boxes.',
        f'**Split (official):** train {tot["train"]:,} / validation {tot["val"]} / test {tot["test"]} pages.',
        f'**Verification benchmark:** {C.n_pairs} English–Hindi and English–Tamil paper pairs (18 questions each), '
        f'{C.n_err} with one injected error of 9 real kinds, {C.n_pairs - C.n_err} clean; half degraded like scans. '
        'A separate development set (48 pairs) was used to build the parser.',
        '**Pre-processing:** Parquet → PNG + COCO JSON, boxes clipped, COCO → YOLO labels, letterbox 640 / 1024 px; '
        'PDF pages rendered at 200 dpi.',
        '**Augmentation:** YOLOv8 mosaic, HSV, scale; Faster R-CNN colour jitter, blur; no mirror flips.',
        '**Challenges:** class imbalance ("paragraph" ≈ 32% of boxes); no public set of erroneous exam papers, hence '
        'the benchmark.',
    ], 1.1, 5.6, size=12)
    caption(s, 'Actual data used', L, 6.65, CW, size=11, italic=False)
    imgs = [(C.FIG / 'bench_en.jpg', 'Benchmark: English source'), (C.FIG / 'bench_hi.jpg', 'Benchmark: Hindi version'),
            (C.ROOT / 'dataset' / 'samples' / 'hindi_newspapers.jpg', 'IndicDLP: Hindi newspaper')]
    for i, (pth, cap) in enumerate(imgs):
        x = L + 0.15 + i * 2.45
        image(s, pth, x, 7.0, height=3.2)
        caption(s, cap, x - 0.2, 10.25, 2.5, size=9)

    # 9. Deep learning model architecture
    s = new()
    title(s, 'Deep Learning Model Architecture')
    bullets(s, [
        f'**Model selected:** DocLayout-YOLO ({A.get("parameters", 0) / 1e6:.1f} M parameters), fine-tuned on the full '
        'IndicDLP by its authors (pretrained, transfer learning); compared with YOLOv8s and Faster R-CNN fine-tuned by us.',
        '**Input / output:** a page image (1024×1024×3 for DocLayout-YOLO, 640×640×3 for YOLOv8s) → N boxes, each with '
        'one of 42 labels and a score.',
        '**Modifications:** COCO heads replaced by 42-class heads (43 with background for Faster R-CNN); losses CIoU + '
        'DFL + BCE (YOLO), RPN + CE + smooth-L1 (Faster R-CNN).',
    ], 1.1, 3, size=12)
    archs = [(f'DocLayout-YOLO, {A.get("parameters", 0) / 1e6:.1f} M (used by IndicDoc Verify)',
              [('Input', '1024² page'), ('Backbone', 'CSP + GL-CRM'), ('Neck', 'PAN, 3 scales'), ('Head', 'YOLOv10, NMS-free'),
               ('Output', 'boxes, 42 classes')]),
             (f'YOLOv8s, {P.get("parameters", 0) / 1e6:.1f} M (fine-tuned by us)',
              [('Input', '640² page'), ('Backbone', 'CSPDarknet C2f + SPPF'), ('Neck', 'PAN-FPN'), ('Head', 'decoupled, anchor-free'),
               ('Output', 'boxes, 42 classes')]),
             (f'Faster R-CNN, {B.get("parameters", 0) / 1e6:.1f} M (baseline)',
              [('Input', '640–1024 px'), ('Backbone', 'ResNet-50'), ('Neck', 'FPN, 5 levels'), ('RPN + RoIAlign', 'proposals'),
               ('Box head', '42 + background')])]
    y = 4.3
    bw, gap = 1.25, 0.26
    for name, parts in archs:
        tb = s.shapes.add_textbox(Inches(L), Inches(y), Inches(CW), Inches(0.3))
        runs(tb.text_frame.paragraphs[0], name, 11.5, bold=True)
        for i, (hd, sub) in enumerate(parts):
            x = L + i * (bw + gap)
            box(s, x, y + 0.38, bw, 0.9, hd, sub, size=10.5,
                fill=RGBColor(0xE6, 0xE6, 0xE6) if 0 < i < 4 else WHITE)
            if i:
                arrow(s, x - gap, y + 0.83, x, y + 0.83)
        y += 1.75
    caption(s, 'Grey: learned feature extraction and detection layers.', L, y - 0.15, CW)

    # 10. Tools, technologies and environment
    s = new()
    title(s, 'Tools, Technologies and Environment')
    table(s, ['Component', 'Details to report'], C.tools, 1.15, [2.1, 5.2], size=11, row_h=0.75, bold_first_col=False)

    # 11. Implementation details
    s = new()
    title(s, 'Implementation Details')
    bullets(s, [
        '**What we implemented:** IndicDoc Verify (ml/verify), the OCR stage with the number pass (ml/ocr), training '
        'and evaluation of two detectors, the benchmark, a Streamlit app and a REST API in Docker.',
        '**Data loading:** IndicDLP Parquet → COCO; PDFs rendered at 200 dpi (PyMuPDF).',
        f'**Model:** pretrained DocLayout-YOLO at its high-recall threshold ({A.get("high_recall_conf")}); YOLOv8s and Faster '
        'R-CNN fine-tuned from COCO weights.',
        '**Loss / optimiser:** YOLOv8s CIoU + DFL + BCE, auto optimiser; Faster R-CNN SGD, lr 0.005, batch 4.',
        '**Inference:** layout → OCR per region → question building → checks → report.',
        '**Checks:** numbering, duplicates, options, marks; translation parity by Needleman–Wunsch alignment; repeats '
        'by character 3-gram similarity; JSON format profiles.',
        '**Interface:** Streamlit "Verify Paper" tab and POST /verify (JSON or HTML report).',
    ], 1.1, 4.9, size=12)
    code = ('# every page: layout (DL) -> OCR with number pass\n'
            'out = predict_single(model, page, "doclayout", device, conf=0.15)\n'
            'doc = extract_document(page, out, "hin", number_pass=True)\n\n'
            '# ml/verify: questions, checks, parity, repeats\n'
            'report = verify(hindi_pages, source_pages=english_pages,\n'
            '                bank=earlier_questions, profile="question_paper")\n'
            'html = to_html(report, {"paper": page_images})')
    code_box(s, code, L, 6.05, CW, 1.75, size=9.5)
    figure(s, C.SHOTS / '07_verify_report.png', L + 1.95, 7.95, width=3.4, cap='Verification report produced by the system')

    # 12. Experimental setup and training
    s = new()
    title(s, 'Experimental Setup and Training')
    bullets(s, [
        '**Dataset sizes:** train ' f'{tot["train"]:,} / validation {tot["val"]} / test {tot["test"]} pages; '
        f'benchmark {C.n_pairs} paper pairs (test) + 48 (development).',
        '**Configuration:** Faster R-CNN batch 4, SGD 0.005, **2 epochs, 76 min**; YOLOv8s 640 px, batch 8, '
        '**12 epochs, 60 min**; Apple M1 Pro GPU (PyTorch MPS).',
        f'**E1 Layout:** 3 detectors on {A.get("num_images")} test pages: mAP@50, mAP@50:95, P / R / F1, per label.',
        '**E2 Robustness:** blur, noise, JPEG, half resolution, faded print.',
        '**E3 Verification:** injected errors, real Tesseract OCR; an error counts only if found at the right question.',
        '**E4 Ablations:** without the number pass; without layout-based options; with a noisy layout.',
        '**Protocol:** parser built on the development set; test set run once.',
    ], 1.1, 5, size=12)
    figure(s, C.FIG / 'fig_training.png', L, 6.5, width=CW,
           cap='Training and validation curves of the two detectors we trained (measured).')

    # 13. Results and performance evaluation
    s = new()
    title(s, 'Results and Performance Evaluation')

    def lrow(n, m):
        rc = m['at_recommended_conf']
        return (n, f3(m['map50']), f3(m['map50_95']), f3(rc['precision']), f3(rc['recall']), f3(rc['f1']))
    caption(s, f'Layout detection ({A.get("num_images")} test pages)', L, 1.05, CW, size=11, italic=False)
    table(s, ['Model', 'mAP@50', 'mAP@50:95', 'Precision', 'Recall', 'F1'],
          [lrow('DocLayout-YOLO (authors)', A), lrow('YOLOv8s (ours)', P), lrow('Faster R-CNN (ours)', B)],
          1.4, [2.3, 1.0, 1.1, 1.0, 0.95, 0.95], size=10.5, row_h=0.32)
    caption(s, 'AP@50 on question-paper labels', L, 2.85, CW, size=11, italic=False)
    table(s, ['Label', 'DocLayout-YOLO', 'YOLOv8s', 'Faster R-CNN'],
          [(k, f3(QAP.get(k, 0)), f3(P.get('per_class_ap50', {}).get(k, 0)), f3(B.get('per_class_ap50', {}).get(k, 0)))
           for k in C.q_labels], 3.2, [2.5, 1.6, 1.6, 1.6], size=10.5, row_h=0.3)
    bullets(s, [
        f'**Verification:** {pct(cq.get("found_at_place", 0))} of injected errors found at the right question on '
        f'born-digital papers, {pct(dq.get("found_at_place", 0))} on scan-degraded pages.',
        f'**False alarms:** {cq.get("false_alarms_per_clean_pair", 0):.1f} per clean pair (born-digital), '
        f'{dq.get("false_alarms_per_clean_pair", 0):.1f} (degraded); each is shown on the page for the vetter.',
        f'**Robust:** {pct(VN.get("overall", {}).get("found_at_place", 0))} still found with a noisy layout; '
        f'{C.n_exact}/{C.n_pairs} papers had every question recovered.',
    ], 5.15, 1.7, size=11.5)
    figure(s, C.FIG / 'fig_verify_by_type.png', L + 1.0, 6.95, width=CW - 2.0,
           cap='Errors found at the right question, per error type.')

    # 14. Comparison with baseline / existing methods
    s = new()
    title(s, 'Comparison with Baseline / Existing Methods')
    caption(s, 'Layout detectors (metric: COCO mAP)', L, 1.05, CW, size=11, italic=False)
    table(s, ['Detector', 'Training data', 'mAP@50', 'mAP@50:95'],
          [('Faster R-CNN (ours, baseline)', '2,293 pages, 76 min', f3(B['map50']), f3(B['map50_95'])),
           ('YOLOv8s (ours)', '2,293 pages, 60 min', f3(P['map50']), f3(P['map50_95'])),
           ('DocLayout-YOLO (authors)', '~96K pages', f3(A['map50']), f3(A['map50_95'])),
           ('Published DocLayout-YOLO [1]', 'full test set', '≈0.735', '0.545')],
          1.4, [2.6, 2.1, 1.3, 1.3], size=10.5, row_h=0.3)
    caption(s, 'Verification ablation (test set)', L, 3.05, CW, size=11, italic=False)
    table(s, ['System', 'Errors found', 'False alarms / clean pair'],
          [(n, pct(r.get('overall', {}).get('found_at_place', 0)), f"{r.get('overall', {}).get('false_alarms_per_clean_pair', 0):.1f}")
           for n, r in (('Full system', V), ('Without number pass', C.V_NP), ('Without layout-based options', C.V_NG),
                        ('Noisy layout', VN)) if r], 3.4, [3.3, 1.8, 2.2], size=10.5, row_h=0.3)
    caption(s, 'Existing approaches', L, 5.0, CW, size=11, italic=False)
    table(s, ['Approach', 'Checks the paper?', 'Compares versions?', 'Indic, offline?'],
          [('Manual vetting (status quo)', 'yes, by hand', 'by hand', 'yes'),
           ('Answer-sheet AI (Eklavvya, GradeLab)', 'no, grades answers', 'no', 'partly'),
           ('MT quality estimation (CometKiwi [10])', 'no structure', 'sentences only', 'cloud / GPU'),
           ('LLM back-translation [9]', 'no structure', 'with human review', 'cloud'),
           ('**IndicDoc Verify (ours)**', '**numbering, options, marks, repeats**', '**question by question**',
            '**12 languages, offline**')],
          5.35, [2.5, 1.75, 1.55, 1.5], size=10, row_h=0.42)
    bullets(s, [f'YOLOv8s beats the Faster R-CNN baseline ({f3(P["map50"])} vs {f3(B["map50"])} mAP@50) at '
                f'{B["inference_time"] / P["inference_time"]:.0f}× the speed; the released full-data model is best and is used.',
                'The number pass is the largest single gain: without it, errors found fall to '
                f'{pct(C.V_NP.get("overall", {}).get("found_at_place", 0))} and false alarms triple.'],
          8.2, 2, size=11.5)

    # 15. Sample outputs and demonstration
    s = new()
    title(s, 'Sample Outputs and Demonstration')
    y = figure(s, C.SHOTS / '08_verify_page.png', L + 0.9, 1.1, width=CW - 1.8,
               cap='**Correct:** a changed number in the Hindi version found and outlined on the page.')
    y = figure(s, C.SHOTS / '09_false_alarm.png', L + 0.9, y + 0.1, width=CW - 1.8,
               cap='**Incorrect:** a false alarm on a degraded scan (OCR lost a digit of an option).')
    figure(s, C.SHOTS / '02_detection_tamil.png', L, y + 0.15, width=3.55,
           cap='Running application: layout regions of a Tamil page', cap_size=9)
    figure(s, C.PLOTS / 'authors_confusion_matrix.png', L + 4.25, y + 0.15, height=2.55,
           cap='Confusion matrix of the layout model', cap_size=9)

    # 16. Challenges and limitations
    s = new()
    title(s, 'Challenges and Limitations')
    bullets(s, [strip(f'**{t}:** {x}') for t, x in (
        ('OCR is the weakest link', 'Indic Tesseract models misread Latin digits ("1947" → "947") and glue option markers '
         'to values; the number pass, layout-based option splitting and page-level digit-script detection were needed '
         'to bring false alarms down.'),
        ('No public data of real errors', 'Exam bodies do not publish erroneous pre-print papers, so we built a benchmark '
         'with realistic injected errors, kept a separate development set, and report clean and scanned results '
         'separately.'),
        ('Meaning is not checked yet', 'Parity catches numbers, options, marks and missing questions without '
         'translation; a pure change of meaning needs a cross-lingual model (hook present, not yet evaluated).'),
        ('Compute and data size', 'Our detectors saw 2.3K pages for about an hour (0.26 vs 0.76 mAP@50 for the released '
         'model), so the product uses the released model; our contribution is the checker built on it.'))],
        1.15, 8, size=13, gap=10)

    # 17. Conclusion and future scope
    s = new()
    title(s, 'Conclusion and Future Scope')
    bullets(s, [
        '**Conclusion**',
        ('**Implemented:** IndicDoc Verify: layout detection (3 detectors compared), OCR for 12 scripts with a number '
         'pass, question parsing, structure / parity / repeat / format checks, report, web app, REST API, Docker.', 1),
        (f'**Results:** {pct(cq.get("found_at_place", 0))} of injected errors found at the right question on born-digital '
         f'papers ({pct(dq.get("found_at_place", 0))} on degraded scans), {cq.get("false_alarms_per_clean_pair", 0):.1f} '
         f'false alarms per clean pair; best layout model {f3(A["map50"])} mAP@50.', 1),
        ('**Objectives:** all four met; our own detectors did not reach the released model, which we report openly.', 1),
        '',
        '**Future scope**',
        ('Pilot with a university exam cell or a state board on real pre-print papers.', 1),
        ('Meaning check with a cross-lingual sentence model; number words and formulas.', 1),
        ('Better Indic OCR behind the same interface; answer-key consistency check.', 1),
        ('More document profiles: bilingual government circulars (Official Languages Act s.3(3)), forms, certificates.', 1),
        ('Deployment: offline licence per exam body (NTA, UPSC, state PSCs, boards, universities).', 1),
    ], 1.1, 9, size=13, gap=6)

    # 18. References
    s = new()
    title(s, 'References')
    bullets(s, [('#', k, strip(r)) for k, r in enumerate(C.refs, 1)], 1.05, 7.3, size=10.5, gap=2)
    bullets(s, ['**Other sources:**', *[(o, 1) for o in C.other]], 8.15, 2.6, size=10.5, gap=1)

    prs.save(OUT)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', str(OUT.parent), str(OUT)],
                   check=True, capture_output=True, timeout=300)
    print(f'{OUT}\n{OUT.with_suffix(".pdf")}  ({len(prs.slides)} slides)')


if __name__ == '__main__':
    main()
